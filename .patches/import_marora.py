"""Import 'Maroara 120 V4.xlsx' into a NEW company through the ERP's own API, then compare results.

Usage:
  python3 import_marora.py <project_dir> <xlsx> <superadmin_email> <superadmin_password> [--dry-copy <dir>]

--dry-copy copies the databases to <dir> first and works on the copies (nothing real changes).
Every call is logged with the sheet/row it came from and whether the system accepted it.
"""
import json
import os
import re
import shutil
import sys
from datetime import date, datetime

args = sys.argv[1:]
proj, xlsx, sa_email, sa_password = args[:4]
dry = args[args.index("--dry-copy") + 1] if "--dry-copy" in args else None
if dry:
    os.makedirs(dry, exist_ok=True)
    for f in ("haven.db", "platform.db"):
        shutil.copy(os.path.join(proj, "db", f), os.path.join(dry, f))
    if os.path.isdir(os.path.join(proj, "db", "tenants")):
        shutil.copytree(os.path.join(proj, "db", "tenants"), os.path.join(dry, "tenants"), dirs_exist_ok=True)
    os.environ["ERP_DB_PATH"] = os.path.join(dry, "haven.db")
    os.environ["ERP_PLATFORM_DB_PATH"] = os.path.join(dry, "platform.db")
    os.environ["ERP_TENANTS_DIR"] = os.path.join(dry, "tenants")
sys.path.insert(0, proj)

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from backend.main import app  # noqa: E402

ADMIN_EMAIL = "marora@havenbuilders.pk"
WANTED_PASSWORD = "Marora-2026"
FALLBACK_PASSWORD = "Estate-Test-2026"
COMPANY = "Marora Residency"

LOG = []


def log(sheet, ref, action, resp=None, expect="ok", note="", ok=None):
    status = resp.status_code if resp is not None else None
    try:
        body = resp.json() if resp is not None else None
    except ValueError:
        body = resp.text[:200] if resp is not None else None
    accepted = status is not None and status < 400
    if ok is None:
        ok = accepted if expect == "ok" else (not accepted)
    detail = body.get("detail") if isinstance(body, dict) and "detail" in body else None
    LOG.append({"sheet": sheet, "ref": ref, "action": action, "status": status, "accepted": accepted,
                "expected": expect, "as_expected": ok, "detail": detail if not accepted else "", "note": note})
    return body if accepted else None


def iso(v):
    """Excel dates arrive as datetime or text like '15-04-2026' / '31-Dec-2025'."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d-%b-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def txt(v):
    return None if v is None else str(v).strip() or None


def floor_no(v):
    s = str(v or "").strip().lower()
    if s.startswith("ground"):
        return 0
    m = re.match(r"(\d+)", s)
    return int(m.group(1)) if m else 0


wb = openpyxl.load_workbook(xlsx, data_only=True)


def rows(sheet, header_row, first_row):
    ws = wb[sheet]
    head = [str(c.value or "").split("\n")[0].strip() for c in ws[header_row]]
    for i, r in enumerate(ws.iter_rows(min_row=first_row, values_only=True), first_row):
        if not any(v not in (None, "") for v in r):
            continue
        yield i, dict(zip(head, r))


client = TestClient(app, raise_server_exceptions=False)
client.__enter__()


def login(email, password):
    client.cookies.clear()
    r = client.post("/api/auth/login", json={"identifier": email, "password": password})
    if r.status_code == 200:
        client.headers["X-CSRF-Token"] = r.json()["csrf_token"]
    return r


# ================================================================== 1. company + admin
S = "Setup"
r = login(sa_email, sa_password)
log(S, "super admin", f"Sign in as {sa_email}", r)
plans = client.get("/api/console/plans").json()
plan = next(p for p in plans if p["code"] == "growth")
existing = [c for c in client.get("/api/console/companies").json() if c["name"] == COMPANY]
if existing:
    log(S, "company", f"Company '{COMPANY}' already exists — stop to avoid a duplicate import", None, ok=False)
    print(json.dumps({"log": LOG}))
    raise SystemExit(0)
r = client.post("/api/console/companies", json={
    "name": COMPANY, "plan_id": plan["id"], "billing_cycle": "monthly", "trial_days": 30, "grace_days": 7,
    "admin_name": "Marora Admin", "admin_email": ADMIN_EMAIL, "city": "Karachi",
    "contact_name": "Ahmed Malik", "notes": "Imported from Maroara 120 V4.xlsx", "seed_sample": False})
created = log(S, "company", f"Create company '{COMPANY}' on the Growth plan with admin {ADMIN_EMAIL}", r)
temp = created["temporary_password"]

r = login(ADMIN_EMAIL, temp)
log(S, "admin", "Admin signs in with the one-time password", r)
me = client.get("/api/projects")
log(S, "admin", "Admin is forced to change password before using the app", me, expect="reject")
r = client.post("/api/auth/change-password", json={"current_password": temp, "new_password": WANTED_PASSWORD})
log(S, "admin", f"Set password '{WANTED_PASSWORD}' (contains the email name 'marora')", r, expect="reject",
    note="Password rule: must not contain the email username")
r = client.post("/api/auth/change-password", json={"current_password": temp, "new_password": FALLBACK_PASSWORD})
log(S, "admin", f"Set password '{FALLBACK_PASSWORD}'", r)
r = login(ADMIN_EMAIL, FALLBACK_PASSWORD)
log(S, "admin", "Admin signs in with the new password", r)
log(S, "admin", "Dashboard of the new company is empty", client.get("/api/dashboard"),
    note=str(client.get("/api/dashboard").json().get("kpi")))

# ================================================================== 2. project, plans, budget
S = "🏗️ Projects"
prj = {}
for i, p in rows(S, 3, 4):
    body = {"name": p["Name"], "location": p["Location"], "city": "Karachi",
            "description": f"{p['Type']} · Manager {p['Manager']} · Excel ID {p['Project ID']}",
            "start_date": iso(p["Start"]), "expected_end_date": iso(p["Target End"]),
            "status": (p["Status"] or "active").lower(), "estimated_cost": int(p["Budget (PKR)"] or 0),
            "current_progress": 0, "number_of_floors": 6, "number_of_units": 11}
    out = log(S, f"row {i}", f"Create project {p['Project ID']} {p['Name']}", client.post("/api/projects", json=body),
              note=f"Start date '{p['Start']}' is text in the sheet; converted to {iso(p['Start'])}")
    if out:
        prj[p["Project ID"]] = out.get("id") or out.get("project", {}).get("id")
PID = prj.get("PRJ-001")

r = client.post("/api/projects", json={"name": "Second Project", "status": "active"})
extra = log(S, "limit", "Growth plan allows more than one project (probe; removed right after)", r)
if extra:
    log(S, "limit", "Delete the probe project", client.delete(f"/api/projects/{extra['id']}"))

S = "📋 Pay Plans"
plans_rows = list(rows(S, 3, 4))
for i, p in plans_rows:
    log(S, f"row {i}", f"Plan {p['Plan ID']} {p['Plan Name']} ({p['Notes']})", None, ok=True,
        note="The system keeps one active plan template per project; Monthly-18 is used by the flats")
r = client.put(f"/api/projects/{PID}/installment-template", json={
    "name": "Monthly-18", "default_booking_bps": 2000,
    "rules": [
        {"label": "Monthly", "trigger_kind": "time", "amount_bps": 8750, "installment_count": 18,
         "start_offset_months": 1, "interval_months": 1},
        {"label": "Possession", "trigger_kind": "construction", "amount_bps": 1250, "milestone_progress": 100},
    ]})
log(S, "PLN-001", "Save Monthly-18 as the project's plan: 20% down, 70% over 18 months, 10% at possession", r)
r = client.put(f"/api/projects/{PID}/installment-template", json={
    "name": "Broken", "default_booking_bps": 2000,
    "rules": [{"label": "Monthly", "trigger_kind": "time", "amount_bps": 7000, "installment_count": 18}]})
log(S, "probe", "Plan whose parts don't add up to 100% is rejected", r, expect="reject")
r = client.post(f"/api/projects/{PID}/installment-template/preview", json={"sale_price": 5670000, "booking_amount": 1134000})
pv = log(S, "PLN-001", "Preview Monthly-18 on a 5,670,000 flat", r)
if pv:
    tot = sum(x["amount"] for x in pv.get("installments", []))
    log(S, "PLN-001", "Preview installments + down payment = sale price", None,
        ok=tot + 1134000 == 5670000, note=f"{len(pv.get('installments', []))} installments totalling {tot:,}")

S = "Budget"
cats = {}
for name in ("Land / plot", "Material", "Labor", "Contractor", "Admin", "Marketing", "Other"):
    out = log(S, name, f"Create budget category '{name}'", client.post("/api/budget/categories", json={"name": name}))
    if out:
        cats[name] = out.get("id")
r = client.post("/api/budget/lines", json={"project_id": PID, "category_id": cats["Other"], "planned_amount": 250000000,
                                           "notes": "Overall project budget from Project Master"})
log(S, "PRJ-001", "Budget line: overall 250,000,000", r)

# ================================================================== 3. flats
S = "🏘️ Flat Register"
units = {}
flat_rows = list(rows(S, 7, 8))
for i, f in flat_rows:
    total = int(f["Total Price [AUTO}"] or 0)
    net = int(f["Net Price [AUTO]"] or 0)
    extra_c, disc = int(f["Extra Charges"] or 0), int(f["Discount"] or 0)
    desc = f"Type {f['Type']} · {f['Area']} sqft @ {f['Price/sqft']}/sqft"
    if extra_c or disc:
        desc += f" · list {total:,} + extra {extra_c:,} − discount {disc:,}"
    if f.get("Notes"):
        desc += f" · {f['Notes']}"
    body = {"project_id": PID, "unit_no": str(f["Flat No."]), "unit_type": "Penthouse" if "anth" in str(f["Type"]) else "Flat",
            "residential_type": f"Type {f['Type']}".strip(), "floor_number": floor_no(f["Floor"]),
            "area_ghaz": round(float(f["Area"]) / 9, 2), "block_tower": "Marora Residency",
            "status": "available", "base_sale_price": net, "description": desc,
            "possession_date": iso(f.get("Possession Date"))}
    out = log(S, f"row {i}", f"Create flat {f['Flat No.']} ({f['Floor'].strip()}, {f['Area']} sqft, {net:,})",
              client.post("/api/units", json=body))
    if out:
        units[f["Flat ID"]] = out.get("id")
r = client.post("/api/units", json={"project_id": PID, "unit_no": "202", "status": "available"})
log(S, "probe", "Creating flat 202 a second time is rejected", r, expect="reject")

# ================================================================== 4. customers
S = "👤 Customers"
customers = {}
for i, c in rows(S, 4, 5):
    body = {"name": txt(c["Customer Name"]), "cnic": txt(c["CNIC"]), "father_name": txt(c["Father/Husband"]),
            "phone": txt(c["Mobile"]), "email": txt(c["Email"]),
            "address": ", ".join(x for x in (txt(c["Address"]), txt(c["City"])) if x),
            "description": f"Excel {c['Customer ID']} · Reg {c['Registration No']} · WhatsApp {c['WhatsApp']}"}
    out = log(S, f"row {i}", f"Register customer {c['Customer ID']} {txt(c['Customer Name'])} (CNIC {c['CNIC']})",
              client.post("/api/customers", json=body),
              expect="ok" if c["Customer ID"] == "C-001" else "reject",
              note="" if c["Customer ID"] == "C-001" else "Same CNIC as C-001 — CNIC must be unique")
    if out:
        customers[c["Customer ID"]] = out.get("id")
r = client.post("/api/customers", json={"name": "No CNIC"})
log(S, "probe", "Customer without CNIC is rejected", r, expect="reject")

# ================================================================== 5. booking for flat 202 + its schedule
S = "🏘️ Flat Register"
f202 = next(f for _, f in flat_rows if f["Flat ID"] == 202)
inst_rows = [r_ for _, r_ in rows("📅 Installments AUTO", 8, 11) if str(r_.get("Inst. ID") or "").startswith("F-202")]
installments = [{"type": (r_["Type"] or "").strip(), "amount": int(r_["Inst. Amount"]), "due_date": iso(r_["Due Date"]),
                 "notes": r_["Inst. ID"], "trigger_kind": "time"} for r_ in inst_rows]
plan_total = sum(x["amount"] for x in installments)
body = {"unit_id": units[202], "project_id": PID, "customer_id": customers["C-001"],
        "booking_date": iso(f202["Booking Date"]), "sale_price": int(f202["Net Price [AUTO]"]),
        "booking_amount": installments[0]["amount"], "payment_mode": "Installments",
        "plan_source": "custom", "installments": installments,
        "notes": f"Plan {f202['Pay Plan']} per Flat Register; schedule from Installments sheet"}
gap = body["sale_price"] - plan_total
booking = log(S, "row 13", f"Book flat 202 to C-001 on {body['booking_date']} for {body['sale_price']:,} "
              f"with the 12-row schedule totalling {plan_total:,}", client.post("/api/bookings", json=body),
              expect="reject" if gap else "ok", note=f"Schedule is {gap:,} short of the sale price")
if not booking and gap:
    body["installments"] = installments + [{
        "type": "Balance (not scheduled in sheet)", "amount": gap, "due_date": iso(f202.get("Possession Date")) or "2027-12-31",
        "notes": "Added on import: the workbook's schedule does not cover the full price", "trigger_kind": "time"}]
    booking = log(S, "row 13", f"Book flat 202 again with a 13th row of {gap:,} due at possession for the unscheduled balance",
                  client.post("/api/bookings", json=body))
BID = (booking or {}).get("id") or (booking or {}).get("booking_id")
r = client.post("/api/bookings", json=body)
log(S, "probe", "Booking flat 202 a second time is rejected", r, expect="reject")

# flat 201: 'Booked' with no customer; flat 2: 'Hold' with no customer
r = client.post(f"/api/units/{units[2]}/holds", json={"notes": "Hold per Flat Register (no customer named)",
                                                      "held_at": iso(next(f for _, f in flat_rows if f['Flat ID'] == 2)["Booking Date"])})
log(S, "row 9", "Put flat 2 on hold without a customer (Flat Register says Hold)", r)
r = client.put(f"/api/units/{units[201]}/status", json={"status": "booked"})
log(S, "row 12", "Mark flat 201 'Booked' with no customer and no booking", r, expect="reject",
    note="Booked/Sold only comes from a real booking, so flat 201 stays Available")
r = client.post("/api/units", json={"project_id": PID, "unit_no": "PROBE-SOLD", "status": "sold"})
log(S, "probe", "Creating a new flat directly as 'sold' is rejected", r, expect="reject")

# customers linked to Available flats (C-002→101, C-003→1, C-004→401/500) — Flat Register says Available
for fid in (1, 101, 401, 500):
    f = next(x for _, x in flat_rows if x["Flat ID"] == fid)
    log(S, f"flat {fid}", f"Flat {fid} is 'Available' but lists customer {f['Customer ID']} — left available (no booking)",
        None, ok=True, note="Customer could not be registered (duplicate CNIC)" if f["Customer ID"] not in customers else "")

S = "📋 Booking Register"
for i, b in rows(S, 4, 5):
    log(S, f"row {i}", f"{b['Booking ID']}: flat {b['Flat No']} to {b['Customer ID']} on {b['Booking Date']} ({b['Payment Plan']})",
        None, ok=True, note="Not imported: Flat Register shows this flat Available and the workbook's totals exclude it")

# ================================================================== 6. customer payments
S = "💳 PAYMENTS Entery"
inst_ids = {}
if BID:
    detail = client.get(f"/api/units/{units[202]}").json()
    for inst in detail.get("installments", []):
        inst_ids[inst.get("notes")] = inst["id"]
for i, p in rows(S, 8, 11):
    if not p.get("Customer ID"):
        if p.get("Amount (PKR)"):
            body = {"narration": f"{p['Payment Type']} (no customer on the sheet)", "amount": int(p["Amount (PKR)"]),
                    "direction": "in", "category": "Membership fee", "payment_method": "Bank", "project_id": PID}
            log(S, f"row {i}", "Membership fee 150,000 with no date", client.post("/api/ledger", json=body), expect="reject",
                note="Date is required")
            body["entry_date"] = date.today().isoformat()
            body["notes"] = "Date missing in sheet — recorded on import date; HBL online transfer"
            log(S, f"row {i}", f"Membership fee 150,000 into bank, dated {body['entry_date']} (import date)",
                client.post("/api/ledger", json=body), note="No customer on the sheet, so it goes to the cashbook")
        continue
    ref = p.get("Installment Ref.")
    body = {"booking_id": BID, "customer_id": customers["C-001"], "amount": int(p["Amount (PKR)"]),
            "payment_date": iso(p["Date"]), "method": p["Payment Mode"], "bank": txt(p["Bank/Account"]),
            "reference_number": txt(p["Cheque/TT No."]), "received_by": txt(p["Received By"]) or "Admin",
            "notes": f"{p['Receipt No.']} · {p['Payment Type']} · {txt(p['Notes']) or ''}".strip(" ·")}
    if ref in inst_ids:
        body["installment_id"] = inst_ids[ref]
    label = f"{p['Receipt No.']}: {int(p['Amount (PKR)']):,} {p['Payment Mode']} on {body['payment_date'] or '(no date)'}"
    if ref:
        label += f" against {ref}"
    r = client.post("/api/payments", json=body)
    expect = "ok"
    note = ""
    if p["Receipt No."] == "R-005":
        expect, note = "reject", "280,000 against F-202-02 whose amount is 150,000"
    if not body["payment_date"]:
        note = "No date on the sheet — the system dates it today"
    out = log(S, f"row {i}", label, r, expect=expect, note=note)
    if not out and r.status_code >= 400 and "installment_id" in body:
        body.pop("installment_id")
        log(S, f"row {i}", f"{p['Receipt No.']} again without a fixed installment (system allocates oldest-first)",
            client.post("/api/payments", json=body))
r = client.post("/api/payments", json={"booking_id": BID, "customer_id": customers["C-001"], "amount": 0})
log(S, "probe", "Zero payment is rejected", r, expect="reject")
r = client.post("/api/payments", json={"booking_id": BID, "customer_id": customers["C-001"], "amount": -5000})
log(S, "probe", "Negative payment is rejected", r, expect="reject")
r = client.post("/api/payments", json={"booking_id": BID, "customer_id": customers["C-001"], "amount": 99_000_000,
                                       "payment_date": date.today().isoformat()})
log(S, "probe", "Payment larger than the whole outstanding balance is rejected", r, expect="reject")
r = client.post("/api/payments", json={"booking_id": BID, "customer_id": customers["C-001"], "amount": 1000,
                                       "payment_date": "2031-01-01"})
log(S, "probe", "Payment dated years in the future is rejected", r, expect="reject")
r = client.post("/api/payments", json={"booking_id": BID, "customer_id": customers["C-001"], "amount": 1000,
                                       "payment_date": "15-07-2026"})
log(S, "probe", "Payment date in dd-mm-yyyy text is rejected", r, expect="reject")
r = client.post("/api/payments", json={"booking_id": 999999, "customer_id": customers["C-001"], "amount": 1000})
log(S, "probe", "Payment on a booking that doesn't exist is rejected", r, expect="reject")

# ================================================================== 7. partners
S = "👥 Partners"
partners = {}
for i, p in rows(S, 7, 8):
    if not str(p.get("Partner ID") or "").startswith("P-"):
        break
    body = {"name": txt(p["Partner Name"]), "cnic": txt(p["CNIC"]), "mobile_number": txt(p["Mobile"]),
            "description": f"Excel {p['Partner ID']} · Father {txt(p['Father/Guardian'])} · {p['Bank']} {p['Bank Account']} · {p['Notes']}",
            "status": (p["Status"] or "active").lower(), "partner_type": "profit_share", "project_id": PID,
            "agreed_amount": int(p["Capital (PKR)"]), "investment_date": iso(p["Join Date"]),
            "profit_share_pct": float(p["Profit %"]) * 100, "profit_share_basis": "net_profit"}
    dup = p["Partner ID"] == "P-003"
    out = log(S, f"row {i}", f"Partner {p['Partner ID']} {txt(p['Partner Name'])}: capital {int(p['Capital (PKR)']):,}, "
              f"{float(p['Profit %']) * 100:.0f}% share", client.post("/api/partners", json=body),
              expect="reject" if dup else "ok",
              note="Same CNIC 42101-0000002-2 as P-002 in the sheet" if dup else "")
    if not out and dup:
        body["cnic"] = None
        body["description"] += " · CNIC left blank on import: sheet repeats P-002's CNIC"
        out = log(S, f"row {i}", f"Partner {p['Partner ID']} again with CNIC left blank", client.post("/api/partners", json=body))
    if out:
        partners[p["Partner ID"]] = out.get("id")
for i, c in rows(S, 14, 15):
    pid = partners.get(c["Partner ID"])
    if not pid:
        log(S, f"row {i}", f"Contribution {c['Txn ID']} skipped — partner not created", None, ok=False)
        continue
    r = client.post(f"/api/partners/{pid}/contribute", json={
        "amount": int(c["Amount (PKR)"]), "contribution_date": iso(c["Date"]),
        "notes": f"{c['Txn ID']} · {c['Type']} · {txt(c['Bank'])} {c['Ref No.']} · {c['Notes']} · source: {txt(c['Source'])}"})
    log(S, f"row {i}", f"{c['Txn ID']}: {int(c['Amount (PKR)']):,} capital from {c['Partner ID']} on {iso(c['Date'])}", r,
        note="Contribution dated a day before the partner's join date" if iso(c["Date"]) < "2026-05-15" else "")

# ================================================================== 8. vendors, materials, purchases
S = "🏪 Vendors"
vendors = {}
for i, v in rows(S, 3, 4):
    body = {"name": txt(v["Vendor Name"]), "category": txt(v["Category"]),
            "contact": f"{v['Contact']} · {v['Phone']} · {v['Email']}", "ntn": txt(v["NTN"]),
            "description": f"Excel {v['Vendor ID']} · credit {v['Credit Days']} days", "status": (v["Status"] or "").lower()}
    dup = v["Vendor ID"] == "V-004"
    out = log(S, f"row {i}", f"Vendor {v['Vendor ID']} {txt(v['Vendor Name'])} ({txt(v['Category'])}, NTN {v['NTN']})",
              client.post("/api/vendors", json=body), expect="reject" if dup else "ok",
              note="Same NTN, contact and email as V-003" if dup else "")
    if not out and dup:
        body["ntn"] = None
        body["description"] += " · NTN left blank on import: sheet repeats V-003's NTN"
        out = log(S, f"row {i}", f"Vendor {v['Vendor ID']} again with NTN left blank", client.post("/api/vendors", json=body))
    if out:
        vendors[v["Vendor ID"]] = out.get("id")

S = "📦 Materials"
materials = {}
for i, m in rows(S, 3, 4):
    body = {"name": m["Name"].strip(), "sku": m["Material ID"], "unit": m["Unit"], "category": m["Category"],
            "project_id": PID, "min_stock": float(m["Reorder Lvl"] or 0), "unit_cost": int(m["Unit Cost"] or 0),
            "notes": f"Usual vendor {m['Vendor ID']}"}
    out = log(S, f"row {i}", f"Material {m['Material ID']} {m['Name'].strip()} ({m['Unit']} @ {int(m['Unit Cost']):,})",
              client.post("/api/inventory", json=body),
              note="Cement priced per '1000 pcs'" if m["Material ID"] == "M-004" else "")
    if out:
        materials[m["Material ID"]] = out

S = "📦 PURCHASES ← Enter Here"
for i, po in rows(S, 5, 6):
    if not po.get("Vendor ID"):
        continue
    body = {"po_no": po["PO No."], "vendor_id": vendors.get(po["Vendor ID"]), "project_id": PID,
            "material": po["Material Name"].strip(), "quantity": str(po["Qty"]), "unit_cost": int(po["Unit Cost"] or 0),
            "total": int(po["Total Cost"] or 0), "order_date": iso(po["Date"]),
            "expected_delivery_date": iso(po["Delivery Date"]), "category": "Material", "budget_category_id": cats.get("Material"),
            "notes": f"Status in sheet: {po['Status']}"}
    out = log(S, f"row {i}", f"{po['PO No.']}: {po['Material Name'].strip()} × {po['Qty']} = {int(po['Total Cost'] or 0):,}",
              client.post("/api/purchase-orders", json=body), expect="reject",
              note="Quantity is 0, so the order is worth nothing")
r = client.post("/api/purchase-orders", json={"vendor_id": vendors.get("V-001"), "project_id": PID, "material": "Probe",
                                               "quantity": "-5", "unit_cost": 1000, "order_date": date.today().isoformat()})
log(S, "probe", "Purchase order with negative quantity is rejected", r, expect="reject")
r = client.post("/api/vendor-payments", json={"vendor_id": vendors.get("V-001"), "amount": 0, "payment_method": "Cash"})
log(S, "probe", "Vendor payment of 0 is rejected", r, expect="reject")

# ================================================================== 9. contractors + site log
S = "👷 Contractors"
contractors = {}
for i, c in rows(S, 3, 4):
    body = {"name": txt(c["Name"]), "specialty": c["Specialty"], "contact": str(c["Contact"]),
            "description": f"Excel {c['Contractor ID']} · retention {float(c['Ret. %']) * 100:.0f}%",
            "status": (c["Status"] or "").lower()}
    out = log(S, f"row {i}", f"Contractor {c['Contractor ID']} {txt(c['Name'])} ({c['Specialty']})",
              client.post("/api/contractors", json=body))
    if out:
        contractors[c["Contractor ID"]] = out.get("id")
for i, sl in rows("🏗️ Site Log", 3, 4):
    cid = contractors.get(sl["Contractor ID"])
    r = client.post(f"/api/contractors/{cid}/assign", json={"project_id": PID, "role": sl["Activity"], "contract_amount": 0,
                                                            "start_date": iso(sl["Date"]), "status": "active",
                                                            "notes": "Contract value 0 in sheet; retention 10%"})
    log("🏗️ Site Log", f"row {i}", f"Assign {sl['Contractor ID']} to the project for {sl['Activity']}", r)
    r = client.post("/api/site-logs", json={"project_id": PID, "log_date": iso(sl["Date"]), "engineer": sl["Approved By"],
                                            "work_done": f"{sl['Activity']} — {sl['Milestone']} ({sl['Contractor ID']})",
                                            "material_used": f"{sl['Material ID']}: {sl['Qty Used']} {sl['Unit']}",
                                            "current_progress": int(round(float(sl["Progress %"]) * 100))})
    log("🏗️ Site Log", f"row {i}", f"Site log {iso(sl['Date'])}: {sl['Milestone']}, progress {float(sl['Progress %']) * 100:.0f}%", r)
    r = client.post(f"/api/contractors/{cid}/pay", json={"amount": 0, "project_id": PID})
    log("🏗️ Site Log", "probe", "Contractor bill/payment of 0 is rejected", r, expect="reject")

# ================================================================== 10. expenses
S = "🧾 EXPENSES ← Enter Here"
for i, e in rows(S, 8, 11):
    if not e.get("Voucher No."):
        continue
    method = "Cash" if str(e["Paid From"]).strip().lower() == "cash" else "Bank"
    body = {"entry_date": iso(e["Date"]), "narration": f"{e['Voucher No.']} · {txt(e['Notes'])} · {txt(e['Vendor/Payee'])}",
            "amount": int(e["Net Amt"]), "direction": "out",
            "category": "Land / plot" if "Plot" in str(e["Notes"]) and e["Sub-Category"].strip() == "Asset" else "Commission"
            if "ommission" in str(e["Sub-Category"]) else (e["Category"] or "Other"),
            "payment_method": method, "project_id": prj.get(e["Project ID"]),
            "notes": f"{e['Category']} / {txt(e['Sub-Category'])} · paid from {txt(e['Paid From'])} · {e['Cheque/TT']} · approved {e['Approved By']}"}
    note = ""
    if e["Voucher No."] == "PV-002":
        note = "Identical to PV-001 (same date, payee, amount and cheque CHQ-5001) — likely a duplicate"
    if not e.get("Project ID"):
        note = "No project in the sheet — company-wide"
    log(S, f"row {i}", f"{e['Voucher No.']}: {int(e['Net Amt']):,} {txt(e['Sub-Category'])} paid from {txt(e['Paid From'])}",
        client.post("/api/ledger", json=body), note=note)
r = client.post("/api/ledger", json={"entry_date": "03-06-2026", "narration": "probe", "amount": 100, "direction": "out"})
log(S, "probe", "Expense dated in dd-mm-yyyy text is rejected", r, expect="reject")
r = client.post("/api/ledger", json={"entry_date": "2026-06-03", "narration": "probe", "amount": 0, "direction": "out"})
log(S, "probe", "Expense of 0 is rejected", r, expect="reject")
r = client.post("/api/ledger", json={"entry_date": "2026-06-03", "narration": "probe", "amount": 10, "direction": "out",
                                     "payment_method": "Crypto"})
log(S, "probe", "Expense paid from an unknown account is rejected", r, expect="reject")

# ================================================================== 11. results vs workbook
get = lambda p: client.get(p).json()
dash = get("/api/dashboard")
ov = get("/api/reports/overview?date_from=2000-01-01&date_to=2099-12-31")
inv = get("/api/reports/inventory")
bs = get("/api/reports/balance-sheet?date_to=2099-12-31")
tb = get("/api/reports/trial-balance?date_to=2099-12-31")
mo = get("/api/reports/monthly")
ex = get("/api/reports/expense-ledger")
ag = get("/api/reports/ageing")
cb = get("/api/reports/customer-balances")
pw = get("/api/reports/project-wise")
audit = get("/api/reports/audit")
acct = {a["code"]: a["debit"] - a["credit"] for a in tb["accounts"]}
revenue = -sum(v for k, v in acct.items() if k.startswith("4"))
expenses = sum(v for k, v in acct.items() if k.startswith("5"))
unit202 = get(f"/api/units/{units[202]}") if units.get(202) else {}
insts = unit202.get("installments", [])

COMPARE = [
    # (what, workbook value, system value, explanation if different)
    ("Total units", 11, inv and sum(p["total"] for p in inv["by_project"]), ""),
    ("Sold units", 1, sum(p["booked"] + p["delivered"] for p in inv["by_project"]), ""),
    ("Available units", 8, sum(p["available"] for p in inv["by_project"]), "Includes flat 201 (see Hold/Booked)"),
    ("Hold/Booked units", 2, sum(p["hold"] for p in inv["by_project"]),
     "Flat 201 cannot be 'Booked' without a customer and booking, so it stays Available"),
    ("Total sales value", 5670000, ov["sales_value"], ""),
    ("Collected from customers (flat 202)", 2430000, cb["totals"]["paid"], ""),
    ("Total collected incl. membership fee", 2580000, cb["totals"]["paid"] + 150000, ""),
    ("Outstanding on flat 202", 3240000, cb["totals"]["outstanding"], ""),
    ("Total expenses", 4060000, ex["total"], ""),
    ("Cash balance", 1835000, acct.get("1010", 0), ""),
    ("Bank balance (workbook excludes partner capital)", -3315000, acct.get("1020", 0),
     "The workbook's Bank Book leaves out the 2,000,000 partner capital that its own Cash Flow sheet includes"),
    ("Cash + bank (workbook Cash Flow closing)", 520000, acct.get("1010", 0) + acct.get("1020", 0), ""),
    ("Partner capital", 2000000, -acct.get("3000", 0), ""),
    ("Net profit", 1610000, revenue - expenses,
     "System also counts the 150,000 membership fee as income; the workbook's P&L leaves it out"),
    ("Vendor payables", 0, -acct.get("2000", 0), ""),
    ("Construction progress %", 15, pw["projects"][0]["progress"] if pw["projects"] else None, ""),
    ("Balance sheet balances", "OUT OF BALANCE", "Balanced" if bs["balanced"] else "Out of balance",
     "The workbook's balance sheet does not balance; the system's does"),
    ("Installments on flat 202", 12, len(insts), "13th row added for the 320,000 the sheet never scheduled"),
    ("Overdue installments (flat 202)", 2, sum(1 for x in insts if str(x.get("status", "")).lower() == "overdue"),
     "Depends on today's date and how payments are allocated"),
    ("Overdue amount", 1150000, ag["total"], "Depends on today's date and how payments are allocated"),
]
print(json.dumps({
    "log": LOG,
    "compare": [{"what": w, "workbook": a, "system": b, "match": a == b, "why": why} for w, a, b, why in COMPARE],
    "installments_202": [{k: x.get(k) for k in ("installment_no", "type", "due_date", "amount", "paid_amount",
                                                 "remaining_amount", "status")} for x in insts],
    "audit_checks": audit["checks"],
    "credentials": {"email": ADMIN_EMAIL, "password": FALLBACK_PASSWORD},
}, default=str))
client.__exit__(None, None, None)
